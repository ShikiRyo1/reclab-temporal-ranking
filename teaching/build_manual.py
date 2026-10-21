"""Build a Chinese study manual with tested native OMML equations."""
from __future__ import annotations
import argparse
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

HERE=Path(__file__).resolve().parent
EQUATIONS=[]
FENCE=chr(96)*3

def el(tag,text=None):
    node=OxmlElement("m:"+tag)
    if text is not None: node.text=text
    return node

def math_run(text):
    node=el("r");prop=el("rPr");style=el("sty");style.set(qn("m:val"),"p")
    prop.append(style);node.append(prop);node.append(el("t",text));return node

SYMBOLS={"theta":"θ","ldots":"…","ell":"ℓ","partial":"∂","top":"⊤","eta":"η","mathcal":"","sigma":"σ","lambda":"λ","Theta":"Θ","Delta":"Δ","alpha":"α","gamma":"γ",
"delta":"δ","epsilon":"ε","pi":"π","tau":"τ","kappa":"κ","nabla":"∇","in":"∈",
"geq":"≥","leq":"≤","approx":"≈","cdot":"·","times":"×","cap":"∩","cup":"∪",
"sum":"∑","bigcup":"⋃","min":"min","max":"max","log":"log","exp":"exp",
"lVert":"‖","rVert":"‖","Vert":"‖","|":"‖","{":"{","}":"}","%":"%","_":"_",
"#":"#","quad":"  ","qquad":"    ",",":" ",";":" ","left":"","right":"","bigl":"","bigr":"","!":"","to":"→"}

class LatexParser:
    """Limited explicit grammar: unsupported syntax stops the build."""
    def __init__(self,source): self.source,self.i=source,0
    def group(self):
        if self.i<len(self.source) and self.source[self.i]=="{":
            self.i+=1;nodes=self.sequence("}")
            if self.i>=len(self.source) or self.source[self.i]!="}":raise ValueError("Unclosed group "+self.source)
            self.i+=1;return nodes
        return [self.atom()]
    def atom(self):
        c=self.source[self.i]
        if c=="{":
            nodes=self.group();node=el("box");child=el("e")
            for n in nodes:child.append(n)
            node.append(child);return node
        if c=="\\":
            self.i+=1;match=re.match(r"[A-Za-z]+|.",self.source[self.i:])
            if not match:raise ValueError("Trailing slash")
            command=match.group();self.i+=len(command)
            if command=="frac":
                node=el("f")
                for tag in ("num","den"):
                    child=el(tag)
                    for n in self.group():child.append(n)
                    node.append(child)
                return node
            if command=="sqrt":
                node=el("rad");pr=el("radPr");hide=el("degHide");hide.set(qn("m:val"),"1")
                pr.append(hide);node.append(pr);node.append(el("deg"));child=el("e")
                for n in self.group():child.append(n)
                node.append(child);return node
            if command=="overline":
                node=el("bar");pr=el("barPr");pos=el("pos");pos.set(qn("m:val"),"top")
                pr.append(pos);node.append(pr);child=el("e")
                for n in self.group():child.append(n)
                node.append(child);return node
            if command=="widetilde":
                node=el("acc");pr=el("accPr");char=el("chr");char.set(qn("m:val"),"~")
                pr.append(char);node.append(pr);child=el("e")
                for n in self.group():child.append(n)
                node.append(child);return node
            if command in ("operatorname","mathbb","mathbf","mathrm","text"):
                nodes=self.group();node=el("box");child=el("e")
                for n in nodes:child.append(n)
                node.append(child);return node
            if command not in SYMBOLS:raise ValueError(f"Unsupported command {command}: {self.source}")
            return math_run(SYMBOLS[command])
        self.i+=1;return math_run("−" if c=="-" else c)
    def sequence(self,stop=None):
        out=[]
        while self.i<len(self.source) and self.source[self.i]!=stop:
            if self.source[self.i] in "\n\r":self.i+=1;continue
            node=self.atom();sub=sup=None
            while self.i<len(self.source) and self.source[self.i] in "_^":
                key=self.source[self.i];self.i+=1;values=self.group()
                if key=="_":sub=values
                else:sup=values
            if sub is not None or sup is not None:
                script=el("sSubSup" if sub is not None and sup is not None else ("sSub" if sub is not None else "sSup"))
                base=el("e");base.append(node);script.append(base)
                for tag,values in (("sub",sub),("sup",sup)):
                    if values is not None:
                        child=el(tag)
                        for n in values:child.append(n)
                        script.append(child)
                node=script
            out.append(node)
        return out

def add_equation(doc,latex):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
    mathp=el("oMathPara");math=el("oMath")
    for n in LatexParser(latex).sequence():math.append(n)
    mathp.append(math);p._p.append(mathp);EQUATIONS.append(latex)

def setfont(run,size=10.5,bold=False):
    run.font.name="Calibri";run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();fonts=rpr.rFonts
    if fonts is None:fonts=OxmlElement("w:rFonts");rpr.insert(0,fonts)
    fonts.set(qn("w:eastAsia"),"Microsoft YaHei");fonts.set(qn("w:ascii"),"Calibri");fonts.set(qn("w:hAnsi"),"Calibri")

def inline(p,text,size=10.5):
    if "\\(" in text:
        for segment in re.split(r"(\\\(.*?\\\))",text):
            if segment.startswith("\\(") and segment.endswith("\\)"):
                math=el("oMath")
                for node in LatexParser(segment[2:-2]).sequence():math.append(node)
                p._p.append(math)
            elif segment:inline(p,segment,size)
        return
    for part in re.split(r"(\x60[^\x60]+\x60|\*\*[^*]+\*\*)",text):
        if not part:continue
        if part.startswith(chr(96)) and part.endswith(chr(96)):
            run=p.add_run(part[1:-1]);setfont(run,size);run.font.name="Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"),"Consolas")
        elif part.startswith("**"):
            run=p.add_run(part[2:-2]);setfont(run,size,True)
        else:
            part=re.sub(r"\[([^\]]+)\]\(([^)]+)\)",r"\1（\2）",part)
            run=p.add_run(part);setfont(run,size)

def configure(doc):
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for name in ("Normal","Body Text"):
        st=doc.styles[name];st.font.name="Calibri";st.font.size=Pt(10.5)
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        st.paragraph_format.line_spacing=1.18;st.paragraph_format.space_after=Pt(7)
    for name,size in (("Title",23),("Heading 1",17),("Heading 2",13)):
        st=doc.styles[name];st.font.name="Calibri";st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(8);st.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    doc.core_properties.title="个性化推荐与排序服务完整学习手册"
    doc.core_properties.subject="MovieLens SAR BPR temporal evaluation and deployment"
    doc.core_properties.author="Yuhui Wu"
    for style in doc.styles:
        for border in list(style._element.iter(qn("w:pBdr"))):border.getparent().remove(border)

def table(doc,lines):
    rows=[[re.sub(r"\x60([^\x60]+)\x60",r"\1",c.strip()) for c in line.strip().strip("|").split("|")] for line in lines]
    rows=[r for r in rows if not all(re.match(r"^:?-+:?$",c) for c in r)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for j,h in enumerate(rows[0]):t.columns[j].width=Cm(17/len(rows[0]));t.rows[0].cells[j].text=h
    for row in rows[1:]:
        for cell,text in zip(t.add_row().cells,row):cell.text=text
    borders=OxmlElement("w:tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        node=OxmlElement("w:"+side);node.set(qn("w:val"),"single");node.set(qn("w:sz"),"4");node.set(qn("w:color"),"D9D9D9");borders.append(node)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr()
        if i==0:rp.append(OxmlElement("w:tblHeader"))
        rp.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();mar=OxmlElement("w:tcMar")
            for side in ("top","left","bottom","right"):
                node=OxmlElement("w:"+side);node.set(qn("w:w"),"90");node.set(qn("w:type"),"dxa");mar.append(node)
            cp.append(mar);shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"DCE6F1" if i==0 else ("F6F8FA" if i%2==0 else "FFFFFF"));cp.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.08
                if i==0:p.paragraph_format.keep_with_next=True
                for r in p.runs:setfont(r,9.5,i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def render_markdown(doc,source,title=True):
    source=re.sub(r"\\\[\s*(.*?)\s*\\\]",lambda m:"$$"+re.sub(r"\s+"," ",m.group(1))+"$$",source,flags=re.S)
    source=re.sub(r"\$\$\s*(.*?)\s*\$\$",lambda m:"$$"+re.sub(r"\s+"," ",m.group(1))+"$$",source,flags=re.S)
    lines=source.splitlines();i=0;first_h1=True
    while i<len(lines):
        line=lines[i].rstrip()
        if not line:i+=1;continue
        if line.startswith(FENCE):
            i+=1;block=[]
            while i<len(lines) and not lines[i].startswith(FENCE):block.append(lines[i]);i+=1
            for text in block:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
                r=p.add_run(text.expandtabs(4));setfont(r,8.4);r.font.name="Consolas";r._element.rPr.rFonts.set(qn("w:ascii"),"Consolas")
            i+=1;continue
        if line.startswith("$$") and line.endswith("$$"):add_equation(doc,line[2:-2]);i+=1;continue
        if line.startswith("|"):
            block=[]
            while i<len(lines) and lines[i].startswith("|"):block.append(lines[i]);i+=1
            table(doc,block);continue
        image=re.match(r"!\[([^\]]*)\]\(([^)]+)\)",line)
        if image:
            path=HERE/image.group(2)
            if not path.exists():raise FileNotFoundError(path)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            with Image.open(path) as im:w,h=im.size
            width=min(6.65,6.0*w/h);p.add_run().add_picture(str(path),width=Inches(width));p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
            cap=doc.add_paragraph();cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;inline(cap,image.group(1),9);i+=1;continue
        if line.startswith("# "):
            if title:doc.add_paragraph(line[2:],style="Title")
            i+=1;continue
        if line.startswith("## "):
            p=doc.add_paragraph(line[3:],style="Heading 1")
            p.paragraph_format.page_break_before=False
            first_h1=False;i+=1;continue
        if line.startswith("### "):doc.add_paragraph(line[4:],style="Heading 2");i+=1;continue
        p=doc.add_paragraph()
        if line.startswith("- "):p.style="List Bullet";line=line[2:]
        inline(p,line);i+=1

def main():
    parser=argparse.ArgumentParser();parser.add_argument("--output",default="个性化推荐与排序服务完整学习手册.docx");args=parser.parse_args()
    files=sorted((HERE/"chapters").glob("*.md"));source="\n\n".join(p.read_text(encoding="utf-8") for p in files)
    doc=Document();configure(doc);render_markdown(doc,source)
    out=HERE/args.output;doc.save(out)
    (HERE/"完整学习手册_源稿.md").write_text(source,encoding="utf-8")
    (HERE/"equations.tex").write_text("\n\n".join("\\[\n"+x+"\n\\]" for x in EQUATIONS),encoding="utf-8")
    report={"output":str(out),"source_files":[str(x.relative_to(HERE)) for x in files],"equations":len(EQUATIONS),"images":len(doc.inline_shapes),"tables":len(doc.tables),"paragraphs":len(doc.paragraphs),"sha256":hashlib.sha256(out.read_bytes()).hexdigest()}
    (HERE/"qa").mkdir(exist_ok=True);(HERE/"qa"/"build_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))
if __name__=="__main__":main()
