"""Build the additive re-audit lesson without touching either original handbook."""
from pathlib import Path
import hashlib,json
from docx import Document
from docx.oxml.ns import qn
import build_manual
from build_manual import render_markdown,EQUATIONS,el
class SupplementLatexParser(build_manual.LatexParser):
    """Add prediction hats without altering the original handbook builder."""
    def atom(self):
        if self.source.startswith('\\hat',self.i):
            self.i+=4
            while self.i<len(self.source) and self.source[self.i].isspace():self.i+=1
            node=el('acc');pr=el('accPr');char=el('chr');char.set(qn('m:val'),'̂')
            pr.append(char);node.append(pr);child=el('e')
            for value in self.group():child.append(value)
            node.append(child);return node
        return super().atom()
build_manual.LatexParser=SupplementLatexParser
HERE=Path(__file__).resolve().parent
parts=[HERE/'环境搭建与工具操作_源稿.md',HERE/'上游Notebook逐行伴读_源稿.md']
missing=[str(p) for p in parts if not p.exists()]
if missing: raise FileNotFoundError('Await complete source parts: '+', '.join(missing))
source='\n\n'.join(p.read_text(encoding='utf-8') for p in parts)
reference=HERE/'个性化推荐与排序服务完整学习手册.docx'
reference_sha=hashlib.sha256(reference.read_bytes()).hexdigest()
doc=Document(reference)
# Preserve the verified page system, real styles and footer; replace body slots only.
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
doc.core_properties.title='环境搭建与上游代码补充手册'
doc.core_properties.subject='P1 teaching re-audit addendum 2026-10-21'
render_markdown(doc,source)
# Source chunks are at most 15 lines. Keep each readable chunk and its label together.
paragraphs=doc.paragraphs
is_code=lambda p:any(r.font.name=='Consolas' for r in p.runs)
for i,p in enumerate(paragraphs):
    if is_code(p):
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=i+1<len(paragraphs) and is_code(paragraphs[i+1])
    if p.text.startswith('代码 L') or p.text=='执行副本新增的 L2 是：':
        p.paragraph_format.keep_with_next=True
    if p.style.name=='Heading 2' and i+1<len(paragraphs) and ('原单元' in p.text or '辅助单元' in p.text):
        paragraphs[i+1].paragraph_format.keep_with_next=True
target=HERE/'环境搭建与上游代码补充手册.docx';doc.save(target)
(HERE/'环境搭建与上游代码补充手册_源稿.md').write_text(source,encoding='utf-8')
equation_source='\n\n'.join('\\[\n'+e+'\n\\]' for e in EQUATIONS)
(HERE/'supplement_equations.tex').write_text(equation_source,encoding='utf-8')
(HERE/'qa/reaudit_20261021/supplement_equations.tex').write_text(equation_source,encoding='utf-8')
report={'sources':[{'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in parts],'reference':reference.name,'reference_sha256':reference_sha,'reference_unchanged':hashlib.sha256(reference.read_bytes()).hexdigest()==reference_sha,'output':target.name,'docx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'equations':len(EQUATIONS),'images':len(doc.inline_shapes),'tables':len(doc.tables)}
(HERE/'qa/reaudit_20261021/supplement_build.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
